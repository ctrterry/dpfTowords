import docx
from docx.shared import Inches
import tkinter as tk
from tkinter import ttk, filedialog, messagebox
from pdf2image import convert_from_path
from PyPDF2 import PdfReader, PdfWriter
import os
import shutil
import tempfile
import threading


def find_poppler_path():
    """Locate Poppler binaries; GUI apps often lack Homebrew in PATH."""
    if shutil.which('pdftoppm'):
        return None

    for path in ('/opt/homebrew/bin', '/usr/local/bin'):
        if os.path.isfile(os.path.join(path, 'pdftoppm')):
            return path

    return None

class PDFConverter:
    def __init__(self):
        self.root = tk.Tk()
        self.root.title("PDF to Word Converter")
        self.root.geometry("800x600") # Tring to 800 x 600
        self.root.configure(bg='#f0f0f0')
        
        # Variables
        self.pdf_path = tk.StringVar()
        self.output_path = tk.StringVar()
        self.conversion_ratio = tk.DoubleVar(value=0.68)
        self.dpi = tk.IntVar(value=300)
        self.poppler_path = find_poppler_path()

        self.setup_ui()
        
    def setup_ui(self):
        # Main frame
        main_frame = ttk.Frame(self.root, padding="20")
        main_frame.pack(fill=tk.BOTH, expand=True)
        
        # File selection
        file_frame = ttk.LabelFrame(main_frame, text="File Selection", padding="10")
        file_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(file_frame, text="PDF File:").pack(side=tk.LEFT, padx=5)
        ttk.Entry(file_frame, textvariable=self.pdf_path, width=50).pack(side=tk.LEFT, padx=5)
        ttk.Button(file_frame, text="Browse", command=self.browse_pdf).pack(side=tk.LEFT, padx=5)
        
        # Output location
        output_frame = ttk.LabelFrame(main_frame, text="Output Location", padding="10")
        output_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(output_frame, text="Output file:").pack(side=tk.LEFT, padx=5)
        ttk.Entry(output_frame, textvariable=self.output_path, width=50).pack(side=tk.LEFT, padx=5)
        ttk.Button(output_frame, text="Rename", command=self.browse_output).pack(side=tk.LEFT, padx=5)
        
        # Settings
        settings_frame = ttk.LabelFrame(main_frame, text="Conversion Settings", padding="10")
        settings_frame.pack(fill=tk.X, pady=5)
        
        ttk.Label(settings_frame, text="Image Ratio:").pack(side=tk.LEFT, padx=5)
        ratio_slider = ttk.Scale(settings_frame, from_=0.1, to=1.0, variable=self.conversion_ratio, 
                 orient=tk.HORIZONTAL, length=200, command=self.update_ratio_label)
        ratio_slider.pack(side=tk.LEFT, padx=5)
        self.ratio_value_label = ttk.Label(settings_frame, text=f"{self.conversion_ratio.get():.2f}")
        self.ratio_value_label.pack(side=tk.LEFT, padx=5)
        
        ttk.Label(settings_frame, text="DPI:").pack(side=tk.LEFT, padx=5)
        ttk.Spinbox(settings_frame, from_=100, to=600, increment=50, 
                   textvariable=self.dpi, width=5).pack(side=tk.LEFT, padx=5)
        
        # Progress
        self.progress_frame = ttk.LabelFrame(main_frame, text="Progress", padding="10")
        self.progress_frame.pack(fill=tk.X, pady=5)
        
        self.progress_bar = ttk.Progressbar(self.progress_frame, mode='determinate')
        self.progress_bar.pack(fill=tk.X, pady=5)
        
        self.status_label = ttk.Label(self.progress_frame, text="Ready")
        self.status_label.pack(pady=5)
        
        # Convert button
        ttk.Button(main_frame, text="Convert", command=self.start_conversion).pack(pady=10)
        
    def _default_output_filename(self, pdf_path):
        stem = os.path.splitext(os.path.basename(pdf_path))[0]
        return f"{stem}_done.docx"

    def _resolve_output_path(self, pdf_path, output_path=None):
        """Always place the output file in the same folder as the input PDF."""
        filename = os.path.basename(output_path) if output_path else ''
        if not filename:
            filename = self._default_output_filename(pdf_path)
        elif not filename.lower().endswith('.docx'):
            filename += '.docx'
        return os.path.join(os.path.dirname(pdf_path), filename)

    def browse_pdf(self):
        filename = filedialog.askopenfilename(
            filetypes=[("PDF files", "*.pdf"), ("All files", "*.*")]
        )
        if filename:
            self.pdf_path.set(filename)
            self.output_path.set(self._resolve_output_path(filename))

    def browse_output(self):
        pdf_path = self.pdf_path.get()
        if not pdf_path:
            messagebox.showerror("Error", "Please select a PDF file first")
            return

        initial_name = os.path.basename(self.output_path.get()) or self._default_output_filename(pdf_path)
        filename = filedialog.asksaveasfilename(
            initialdir=os.path.dirname(pdf_path),
            initialfile=initial_name,
            defaultextension=".docx",
            filetypes=[("Word files", "*.docx"), ("All files", "*.*")]
        )
        if filename:
            self.output_path.set(self._resolve_output_path(pdf_path, filename))
            
    def convert_pdf_to_word(self):
        try:
            pdf_path = self.pdf_path.get()

            if not pdf_path:
                messagebox.showerror("Error", "Please select a PDF file")
                return

            output_path = self._resolve_output_path(pdf_path, self.output_path.get())
            self.output_path.set(output_path)
                
            # Read PDF
            pdf_reader = PdfReader(pdf_path)
            total_pages = len(pdf_reader.pages)
            
            # Create Word document
            doc = docx.Document()
            
            # Create temporary directory
            with tempfile.TemporaryDirectory() as temp_dir:
                for i in range(total_pages):
                    # Update progress
                    progress = (i + 1) / total_pages * 100
                    self.progress_bar['value'] = progress
                    self.status_label['text'] = f"Processing page {i+1} of {total_pages}"
                    self.root.update()
                    
                    # Convert page to image
                    page = pdf_reader.pages[i]
                    pdf_writer = PdfWriter()
                    pdf_writer.add_page(page)
                    
                    temp_pdf = os.path.join(temp_dir, 'temp.pdf')
                    with open(temp_pdf, 'wb') as f:
                        pdf_writer.write(f)
                    
                    convert_kwargs = {'dpi': self.dpi.get()}
                    if self.poppler_path:
                        convert_kwargs['poppler_path'] = self.poppler_path
                    images = convert_from_path(temp_pdf, **convert_kwargs)
                    
                    # Add image to Word document
                    for img in images:
                        temp_img = os.path.join(temp_dir, 'temp.jpg')
                        img.save(temp_img)
                        
                        p = doc.add_paragraph()
                        run = p.add_run()
                        run.add_break()
                        doc.add_picture(
                            temp_img,
                            width=Inches(self.conversion_ratio.get() * 8.5)
                        )
                        doc.add_page_break()
            
            # Save document
            doc.save(output_path)
            
            self.status_label['text'] = "Conversion completed successfully!"
            messagebox.showinfo("Success", "PDF has been converted to Word successfully!")
            
        except Exception as e:
            self.status_label['text'] = f"Error: {str(e)}"
            messagebox.showerror("Error", f"An error occurred: {str(e)}")
            
        finally:
            self.progress_bar['value'] = 0
            
    def start_conversion(self):
        # Start conversion in a separate thread
        thread = threading.Thread(target=self.convert_pdf_to_word)
        thread.daemon = True
        thread.start()
        
    def run(self):
        self.root.mainloop()

    def update_ratio_label(self, event=None):
        self.ratio_value_label.config(text=f"{self.conversion_ratio.get():.2f}")

if __name__ == '__main__':
    app = PDFConverter()
    app.run() 